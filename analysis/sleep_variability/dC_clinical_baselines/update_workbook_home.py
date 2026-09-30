"""
Update COMPARISON_T90_vs_TST.xlsx in place for the home-sleep extension.

  1. Back up the workbook first as COMPARISON_T90_vs_TST_pre_home.xlsx (never overwrites an
     existing backup, so the pre-edit state survives reruns).
  2. Relabel: every column header and sheet name that says "TST" or "total sleep time" gets
     "(PSG)" appended after the token, so every sleep-duration column says which kind it is.
     Data cells are untouched.
  3. Add the "Home sleep block" sheet. Every number on it is read from provenance_home.json,
     written by run_home_baselines.py. Nothing is retyped.
  4. Update the README sheet with the home-sleep definition, source, and subset design.
  5. Verify: reload the saved file, rebuild the expected content by applying the same edit
     plan to the backup snapshot, and require every cell to match. Any drift outside the
     planned edits fails the run.
"""
from __future__ import annotations

import json
import os
import re
import shutil
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(HERE, "COMPARISON_T90_vs_TST.xlsx")
BACKUP = os.path.join(HERE, "COMPARISON_T90_vs_TST_pre_home.xlsx")
PROV = json.load(open(os.path.join(HERE, "provenance_home.json")))

RUNG_LABEL = {
    "H1_age_sex": "age + sex",
    "H2_age_sex_AHI": "age + sex + AHI",
    "H3_age_sex_race_smoke": "age + sex + race + smoking",
    "H4_clinical_AHI_SE": "age + sex + race + smoking + AHI + sleep efficiency",
}
SUM = {r["model"]: r for r in PROV["rung_summary"]}
N_COV = PROV["exposure"]["coverage_n"]
N_COHORT = PROV["exposure"]["cohort_n"]
PCT_COV = PROV["exposure"]["coverage_pct"]
N_ROWS = PROV["analysis_rows"]["n"]
BY_SITE = PROV["analysis_rows"]["by_site"]
N_SURV = PROV["event_floor"]["n_survivors_of_48"]


# ------------------------------------------------------------------ the relabel rule
def relabel(text):
    """Append (PSG) after every TST / total sleep time token, once."""
    if not isinstance(text, str):
        return text
    t = re.sub(r"total sleep time(?! \(PSG\))", "total sleep time (PSG)", text)
    t = re.sub(r"\bTST\b(?! \(PSG\))", "TST (PSG)", t)
    return t


SHEET_RENAME = {"TST ladder detail": relabel("TST ladder detail")}

# README value edits, exact old text so nothing is edited by accident
README_EDITS = {
    ("A", "T90 versus total sleep time"):
        "T90 versus total sleep time (PSG) versus habitual sleep (home)",
    ("B", "How much held-out discrimination each exposure adds on top of progressively "
          "richer clinical baselines, across the paper's 48 outcomes. Supplementary side "
          "analysis. Nothing here enters the manuscript."):
        "How much held-out discrimination each exposure adds on top of progressively "
        "richer clinical baselines, across the paper's 48 outcomes. Exposures: T90, total "
        "sleep time (PSG), and home-reported habitual sleep (the Home sleep block sheet). "
        "Supplementary side analysis. Nothing here enters the manuscript.",
    ("A", "  total sleep time"): "  total sleep time (PSG)",
    ("B", "The T90 ladder's rich rungs hold total sleep time as a covariate. That covariate "
          "cannot survive when total sleep time is the exposure."):
        "The T90 ladder's rich rungs hold total sleep time (PSG) as a covariate. That "
        "covariate cannot survive when total sleep time (PSG) is the exposure.",
    ("A", "  TST ladder detail"): "  TST (PSG) ladder detail",
    ("B", "every model in the total sleep time run, including the age-spline drop check"):
        "every model in the total sleep time (PSG) run, including the age-spline drop check",
    ("B", "the table, 4 decimals as requested. Note that the total sleep time dC values sit "
          "around -0.001, so 4 decimals rounds -0.00096 to -0.0010 and -0.00041 to -0.0004. "
          "Nothing is lost, the Full precision sheet and results_tst.csv carry every digit."):
        "the table, 4 decimals as requested. Note that the total sleep time (PSG) dC values "
        "sit around -0.001, so 4 decimals rounds -0.00096 to -0.0010 and -0.00041 to "
        "-0.0004. Nothing is lost, the Full precision sheet and results_tst.csv carry every "
        "digit.",
    ("B", "results.csv, results_tst.csv, per_outcome.csv, per_outcome_tst.csv, "
          "REPORT_TST.txt, provenance_tst.json"):
        "results.csv, results_tst.csv, per_outcome.csv, per_outcome_tst.csv, "
        "REPORT_TST.txt, provenance_tst.json, results_home.csv, provenance_home.json, "
        "REPORT.txt (HOME SLEEP section)",
    ("B", "run_tst_baselines.py then make_comparison.py"):
        "run_tst_baselines.py then make_comparison.py. Home sleep block by "
        "run_home_baselines.py and update_workbook_home.py.",
}

# rows inserted AFTER the row whose column A equals the anchor text
HOME_BLOCK_ROWS = [
    ("", ""),
    ("HOME-REPORTED HABITUAL SLEEP", "the third exposure, added 2026-08-20"),
    ("  definition",
     "Per-patient habitual home sleep duration in hours from the v2-clean clinical-note "
     "extraction. The value is the 'hours' column of "
     "habitual_sleep_v2/habitual_per_patient_v2.parquet, the exact per-patient value the "
     "paper's banded habitual analysis consumes (habitual_outcomes_v2/build.py merges it "
     "1:1 onto the cohort and uses hours.notna() as the usable filter)."),
    ("  transform",
     "within-site rank inverse normal, ONE column, per 1 SD, recomputed inside the "
     "analysis rows, exactly like the other two exposures"),
    ("  coverage",
     f"available for {N_COV:,} of {N_COHORT:,} ({PCT_COV}% of the cohort). Analysis rows "
     f"I0002 {BY_SITE.get('I0002', 0):,}, I0006 {BY_SITE.get('I0006', 0):,}."),
    ("  subset design",
     f"mirrors the BMI block: every model in the Home sleep block sheet is fitted INSIDE "
     f"the home-covered rows, and each rung carries its OWN matched baseline refitted on "
     f"those rows. The analysis rows are the {N_ROWS:,} home-covered rows also complete on "
     f"TST (PSG) and sleep efficiency ({N_COV - N_ROWS} rows dropped), so T90, total sleep "
     f"time (PSG) and habitual sleep (home) are evaluated on identical patients with an "
     f"identical outcome set within each rung. Never compare these C values with the "
     f"full-cohort or BMI sheets."),
    ("  event floor",
     f"the paper's per-fold floor (a fold is dropped when either side has under 30 events) "
     f"means an outcome survives only when both hospitals contribute at least 30 events, "
     f"so at least 60 events in the subset. {N_SURV} of 48 outcomes survive, the same set "
     f"for all three exposures within each rung."),
]
HOME_BLOCK_ANCHOR_A_PREFIX = "BLOCKS"

SHEETS_ROW = ("  Home sleep block",
              "the three-exposure ladder on the home-covered rows, one row per rung, "
              "habitual sleep (home) beside T90 and total sleep time (PSG)")
SHEETS_ANCHOR_A = "  TST (PSG) ladder detail"      # after the A-column edit is applied


# ------------------------------------------------------------------ expected content
def snapshot(path):
    wb = load_workbook(path, data_only=False)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        out[name] = [[c.value for c in row] for row in ws.iter_rows()]
    wb.close()
    return out


def expected_readme(rows):
    """Apply the README edit plan to the backup rows, return the expected rows."""
    new, n_hit = [], 0
    for row in rows:
        a = row[0] if len(row) > 0 else None
        b = row[1] if len(row) > 1 else None
        if ("A", a) in README_EDITS:
            a = README_EDITS[("A", a)]
            n_hit += 1
        if ("B", b) in README_EDITS:
            b = README_EDITS[("B", b)]
            n_hit += 1
        new.append([a, b])
    assert n_hit == len(README_EDITS), (
        f"only {n_hit} of {len(README_EDITS)} planned README edits matched the workbook, "
        f"a source text must have drifted, stop")
    # insert the home block after the BLOCKS row
    ai = [i for i, r in enumerate(new) if r[0] == HOME_BLOCK_ANCHOR_A_PREFIX]
    assert len(ai) == 1, f"BLOCKS anchor found {len(ai)} times"
    at = ai[0] + 1
    new = new[:at] + [[a, b] for a, b in HOME_BLOCK_ROWS] + new[at:]
    # insert the Home sleep block sheet row after the renamed ladder-detail sheet row
    si = [i for i, r in enumerate(new) if r[0] == SHEETS_ANCHOR_A]
    assert len(si) == 1, f"sheets anchor found {len(si)} times"
    new = new[:si[0] + 1] + [[SHEETS_ROW[0], SHEETS_ROW[1]]] + new[si[0] + 1:]
    return new


def home_sheet_rows():
    """The new sheet's full content, straight from provenance_home.json."""
    hdr = ["Baseline model (rung)", "Baseline C", "dC T90", "dC total sleep time (PSG)",
           "dC habitual sleep (home)", "outcomes analyzed (of 48)", "n rows"]
    rows = [hdr]
    for m, lab in RUNG_LABEL.items():
        r = SUM[m]
        rows.append([lab, r["baseline_C"], r["dC_t90"], r["dC_tst"], r["dC_home"],
                     r["n_outcomes"], r["n_rows"]])
    rows.append([None] * 7)
    rows.append([f"Home sleep from clinical-note extraction (v2), available for "
                 f"{N_COV:,} of {N_COHORT:,}. All three exposures evaluated on identical "
                 f"patients.", None, None, None, None, None, None])
    rows.append([f"Each rung is refitted with its OWN matched baseline on the same "
                 f"{N_ROWS:,} home-covered rows ({N_COV - N_ROWS} of the {N_COV:,} lack "
                 f"TST (PSG) or sleep efficiency and are dropped so the three exposures "
                 f"sit on identical people). The outcome set is identical across the "
                 f"three exposures within each rung under the paper's event floor (30 "
                 f"events per hospital, so 60 in the subset). Never compare these C "
                 f"values with the full-cohort or BMI sheets.",
                 None, None, None, None, None, None])
    rows.append([None] * 7)
    rows.append(["Supplement, same rows and outcome sets, full detail",
                 None, None, None, None, None, None])
    shdr = ["Baseline model (rung)", "C with habitual sleep (home)",
            "rel gain T90", "rel gain total sleep time (PSG)",
            "rel gain habitual sleep (home)", "n outcomes home dC>0", None]
    rows.append(shdr)
    for m, lab in RUNG_LABEL.items():
        r = SUM[m]
        rows.append([lab, r["C_with_home"], r["rel_gain_t90"], r["rel_gain_tst"],
                     r["rel_gain_home"], r["npos_home"], None])
    return rows


def expected_workbook(back):
    """Sheet name -> expected rows, from the backup snapshot plus the edit plan."""
    exp = {}
    order = []
    for name, rows in back.items():
        new_name = SHEET_RENAME.get(name, name)
        if name == "README":
            exp[new_name] = expected_readme(rows)
        else:
            hdr = [relabel(v) for v in rows[0]]
            exp[new_name] = [hdr] + [list(r) for r in rows[1:]]
        order.append(new_name)
    # the new sheet sits right before the ladder-detail sheet
    at = order.index(SHEET_RENAME["TST ladder detail"])
    order.insert(at, "Home sleep block")
    exp["Home sleep block"] = home_sheet_rows()
    return exp, order


# ------------------------------------------------------------------ write
def main():
    if not os.path.exists(BACKUP):
        shutil.copy2(XLSX, BACKUP)
        print(f"backed up -> {BACKUP}")
    else:
        print(f"backup already exists, kept untouched: {BACKUP}")

    back = snapshot(BACKUP)
    assert list(back) == ["README", "Comparison", "Full precision", "Per-organ dC",
                          "TST ladder detail"], list(back)
    exp, order = expected_workbook(back)

    wb = load_workbook(XLSX)
    assert wb.sheetnames == list(back), (
        f"workbook drifted from the backup already: {wb.sheetnames}")

    # 1. sheet rename
    for old, new in SHEET_RENAME.items():
        wb[old].title = new

    # 2. header relabels on the table sheets
    n_hdr = 0
    for name in ["Comparison", "Full precision", "Per-organ dC",
                 SHEET_RENAME["TST ladder detail"]]:
        ws = wb[name]
        for c in ws[1]:
            nv = relabel(c.value)
            if nv != c.value:
                c.value = nv
                n_hdr += 1
    print(f"relabeled {n_hdr} column headers")

    # 3. README: value edits then row insertions (bottom-most anchor first)
    ws = wb["README"]
    n_ed = 0
    for row in ws.iter_rows(min_col=1, max_col=2):
        for c, col in zip(row, ("A", "B")):
            if (col, c.value) in README_EDITS:
                c.value = README_EDITS[(col, c.value)]
                n_ed += 1
    assert n_ed == len(README_EDITS), (
        f"only {n_ed} of {len(README_EDITS)} planned README edits matched, stop")
    print(f"edited {n_ed} README cells")

    def find_row(ws, text):
        hits = [c.row for c in ws["A"] if c.value == text]
        assert len(hits) == 1, f"anchor {text!r} found {len(hits)} times"
        return hits[0]

    key_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    def insert_rows(ws, at, rows):
        ws.insert_rows(at, amount=len(rows))
        for i, (a, b) in enumerate(rows):
            ws.cell(row=at + i, column=1, value=(a if a != "" else None)).font = key_font
            ws.cell(row=at + i, column=2, value=(b if b != "" else None)).alignment = wrap

    r_sheets = find_row(ws, SHEETS_ANCHOR_A)
    insert_rows(ws, r_sheets + 1, [SHEETS_ROW])
    r_blocks = find_row(ws, HOME_BLOCK_ANCHOR_A_PREFIX)
    insert_rows(ws, r_blocks + 1, HOME_BLOCK_ROWS)

    # 4. the new sheet
    at = wb.sheetnames.index(SHEET_RENAME["TST ladder detail"])
    nws = wb.create_sheet("Home sleep block", at)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    rows = exp["Home sleep block"]
    numfmt = {2: "0.0000", 3: "0.00000", 4: "0.00000", 5: "0.00000", 6: "0", 7: "#,##0"}
    sup_numfmt = {2: "0.0000", 3: "0.00%", 4: "0.00%", 5: "0.00%", 6: "0"}
    sup_hdr_row = 1 + len(RUNG_LABEL) + 4 + 1   # header + rungs + blank + 2 notes + blank + title
    for ri, row in enumerate(rows, start=1):
        for ci, v in enumerate(row, start=1):
            cell = nws.cell(row=ri, column=ci, value=v)
            in_main = 2 <= ri <= 1 + len(RUNG_LABEL)
            in_sup = ri > sup_hdr_row + 1
            if ri == 1 or ri == sup_hdr_row + 1:
                if v is not None:
                    cell.font, cell.fill, cell.border = hdr_font, hdr_fill, border
                    cell.alignment = Alignment(horizontal="center", vertical="center",
                                               wrap_text=True)
            elif in_main or in_sup:
                cell.border = border
                fmt = (numfmt if in_main else sup_numfmt).get(ci)
                if fmt and v is not None and not isinstance(v, str):
                    cell.number_format = fmt
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(vertical="top", wrap_text=(ci == 1))
            else:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = {1: 52, 2: 14, 3: 12, 4: 22, 5: 22, 6: 22, 7: 10}
    for ci, w in widths.items():
        nws.column_dimensions[get_column_letter(ci)].width = w
    nws.row_dimensions[1].height = 34
    nws.freeze_panes = "A2"

    wb.save(XLSX)
    print(f"saved {XLSX}")

    # 5. verify: reload and require every cell to equal the expected content
    got = snapshot(XLSX)
    assert list(got) == order, f"sheet order/names differ: {list(got)} vs {order}"
    bad = []
    for name in order:
        e, g = exp[name], got[name]
        nr = max(len(e), len(g))
        for ri in range(nr):
            er = e[ri] if ri < len(e) else []
            gr = g[ri] if ri < len(g) else []
            nc = max(len(er), len(gr))
            for ci in range(nc):
                ev = er[ci] if ci < len(er) else None
                gv = gr[ci] if ci < len(gr) else None
                if ev == "":
                    ev = None
                if gv == "":
                    gv = None
                if isinstance(ev, float) and isinstance(gv, float):
                    same = (ev == gv) or (abs(ev - gv) < 1e-15)
                else:
                    same = ev == gv
                if not same:
                    bad.append((name, ri + 1, ci + 1, ev, gv))
    if bad:
        for b in bad[:30]:
            print("MISMATCH", b)
        sys.exit(f"verification FAILED, {len(bad)} cells differ from the edit plan")
    n_cells = sum(sum(len(r) for r in rows) for rows in got.values())
    print(f"verification PASS: {len(order)} sheets, every cell matches the backup plus the "
          f"planned edits ({n_cells} cells checked)")


if __name__ == "__main__":
    main()
